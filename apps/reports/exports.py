from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

PROJECT_HEADERS = ["Proyecto", "Estado", "Tareas", "Hechas", "Avance %", "Atrasadas", "Horas estimadas", "Horas registradas"]
EMPLOYEE_HEADERS = ["Persona", "Asignadas", "Hechas", "Pendientes", "Atrasadas", "A tiempo", "Con retraso", "Puntualidad %", "Horas registradas"]


def _dash(value):
    return "—" if value is None else value


def _summary_rows(report):
    s = report["summary"]
    return [
        ("Alcance", report["scope"]),
        ("Generado", report["generated_at"].strftime("%d/%m/%Y %H:%M")),
        ("Tareas totales", s["total"]),
        ("Tareas completadas", s["done"]),
        ("Tareas pendientes", s["pending"]),
        ("Avance global %", s["progress"]),
        ("Tareas atrasadas", s["overdue"]),
        ("Entregas a tiempo", s["on_time"]),
        ("Entregas con retraso", s["late"]),
        ("Puntualidad %", _dash(s["punctuality"])),
        ("Horas estimadas", s["estimated"]),
        ("Horas registradas", s["logged"]),
    ]


def _project_rows(report):
    return [
        [p["name"], p["status"], p["total"], p["done"], p["progress"], p["overdue"], p["estimated"], p["logged"]]
        for p in report["projects"]
    ]


def _employee_rows(report):
    return [
        [e["name"], e["assigned"], e["done"], e["pending"], e["overdue"],
         e["on_time"], e["late"], _dash(e["punctuality"]), e["hours"]]
        for e in report["employees"]
    ]


# ---------------- Excel ----------------

HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _fill_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for row in rows:
        ws.append(list(row))
    # Texto que empieza con "=" se guarda como texto, nunca como fórmula
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    for i, column in enumerate(ws.columns, start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 12), 50)


def build_excel(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    _fill_sheet(ws, ["Indicador", "Valor"], _summary_rows(report))
    _fill_sheet(wb.create_sheet("Proyectos"), PROJECT_HEADERS, _project_rows(report))
    _fill_sheet(wb.create_sheet("Equipo"), EMPLOYEE_HEADERS, _employee_rows(report))
    _fill_sheet(wb.create_sheet("Horas por semana"), ["Semana (lunes)", "Horas"], report["weeks"])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------- PDF ----------------

def _table(headers, rows, total_width, fractions, cell_style, head_style):
    data = [[Paragraph(escape(str(h)), head_style) for h in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(v)), cell_style) for v in row])
    table = Table(data, colWidths=[total_width * f for f in fractions], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return table


def build_pdf(report):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), title="Reporte de productividad",
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    width = doc.width

    story = [
        Paragraph("Reporte de productividad", styles["Title"]),
        Paragraph(
            f"{escape(report['scope'])} · generado el {report['generated_at']:%d/%m/%Y %H:%M}",
            styles["Normal"],
        ),
        Spacer(1, 12),
        Paragraph("Resumen", styles["Heading2"]),
        _table(["Indicador", "Valor"], _summary_rows(report)[2:], width, [0.3, 0.2], cell, head),
        Spacer(1, 12),
        Paragraph("Proyectos", styles["Heading2"]),
    ]
    projects = _project_rows(report)
    story.append(
        _table(PROJECT_HEADERS, projects, width, [0.26, 0.12, 0.08, 0.08, 0.09, 0.10, 0.13, 0.14], cell, head)
        if projects else Paragraph("Sin proyectos para mostrar.", styles["Normal"])
    )

    story += [Spacer(1, 12), Paragraph("Rendimiento por persona", styles["Heading2"])]
    employees = _employee_rows(report)
    story.append(
        _table(EMPLOYEE_HEADERS, employees, width, [0.23, 0.09, 0.09, 0.10, 0.10, 0.09, 0.10, 0.10, 0.10], cell, head)
        if employees else Paragraph("Aún no hay tareas asignadas ni horas registradas.", styles["Normal"])
    )

    story += [
        Spacer(1, 12),
        Paragraph("Horas registradas por semana", styles["Heading2"]),
        _table(["Semana (lunes)", "Horas"], report["weeks"], width, [0.15, 0.15], cell, head),
    ]
    doc.build(story)
    return buffer.getvalue()